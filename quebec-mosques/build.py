"""Merge regional mosque search results into one deduplicated CSV and XLSX."""
import csv
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
REGION_FILES = {
    "Montréal (centre/ouest)": "mtl_central_west.json",
    "Montréal (nord/est)": "mtl_north_east.json",
    "Laval, Lanaudière, Laurentides": "laval_north.json",
    "Montérégie (Rive-Sud)": "south_shore.json",
    "Reste du Québec": "rest_of_qc.json",
}
FIELDS = ["region", "name", "type", "address", "city", "postal_code",
          "phone", "website", "confidence", "source"]
EXCLUDE = [r"makkah", r"mecque", r"islamic (centre|center) of quebec", r"\bicq\b",
           r"centre islamique du qu[eé]bec", r"madani"]


def norm(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def street_key(addr):
    m = re.match(r"\s*(\d+)[\s,-]*(.*)", addr or "")
    if not m:
        return ""
    return m.group(1) + norm(m.group(2))[:6]


def rank(row):
    return ({"high": 3, "medium": 2, "low": 1}.get(row["confidence"], 0),
            sum(bool(row[f]) for f in FIELDS))


def main(src_dir):
    rows = []
    for region, fname in REGION_FILES.items():
        path = Path(src_dir) / fname
        if not path.exists():
            print(f"missing {fname}", file=sys.stderr)
            continue
        for r in json.loads(path.read_text()):
            row = {f: str(r.get(f, "") or "").strip() for f in FIELDS}
            row["region"] = region
            row["postal_code"] = row["postal_code"].upper()
            rows.append(row)

    rows = [r for r in rows if not any(re.search(p, r["name"], re.I) for p in EXCLUDE)]

    kept = {}
    for r in rows:
        key = street_key(r["address"]) or norm(r["name"])
        if key not in kept or rank(r) > rank(kept[key]):
            kept[key] = r
    out = sorted(kept.values(), key=lambda r: (r["region"], norm(r["city"]), norm(r["name"])))

    with open(HERE / "quebec_mosques.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(out)

    wb = Workbook()
    ws = wb.active
    ws.title = "Mosquées QC"
    ws.append(FIELDS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E3D")
    for r in out:
        ws.append([r[f] for f in FIELDS])
    widths = {"region": 26, "name": 42, "type": 14, "address": 38, "city": 28,
              "postal_code": 11, "phone": 16, "website": 34, "confidence": 11, "source": 50}
    for i, f in enumerate(FIELDS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[f]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(HERE / "quebec_mosques.xlsx")

    print(f"{len(rows)} raw rows, {len(out)} after dedup")
    for region in REGION_FILES:
        print(f"  {region}: {sum(r['region'] == region for r in out)}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else HERE / "raw")
